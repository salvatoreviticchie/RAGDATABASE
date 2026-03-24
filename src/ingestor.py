from __future__ import annotations

import base64
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openai import OpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings, get_index, get_pc
from .utils import make_vector_id, clean_text

# 4 chars ≈ 1 token; convert token counts to character counts
_CHUNK_SIZE_CHARS = settings.chunk_size * 4
_CHUNK_OVERLAP_CHARS = settings.chunk_overlap * 4

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=_CHUNK_SIZE_CHARS,
    chunk_overlap=_CHUNK_OVERLAP_CHARS,
    separators=["\n\n", "\n", ". ", " ", ""],
)

# Supported file types
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
TEXT_EXTENSIONS   = {".txt"}
DOCX_EXTENSIONS   = {".docx"}
PDF_EXTENSIONS    = {".pdf"}
ALL_EXTENSIONS    = PDF_EXTENSIONS | DOCX_EXTENSIONS | TEXT_EXTENSIONS | IMAGE_EXTENSIONS

_openrouter = OpenAI(
    api_key=settings.openrouter_api_key,
    base_url="https://openrouter.ai/api/v1",
)


# ── Extractors ────────────────────────────────────────────────────────────────

def _extract_pdf(file_path: str) -> list[tuple[int, str]]:
    """Return list of (page_number, text) from a PDF."""
    doc = fitz.open(file_path)
    pages = []
    for i, page in enumerate(doc):
        text = clean_text(page.get_text("text"))
        if text:
            pages.append((i + 1, text))
    return pages


def _extract_docx(file_path: str) -> list[tuple[int, str]]:
    """Extract text from a .docx file. Returns single 'page' of full text."""
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = clean_text("\n".join(paragraphs))
    return [(0, full_text)] if full_text else []


def _extract_txt(file_path: str) -> list[tuple[int, str]]:
    """Extract text from a plain text file."""
    with open(file_path, encoding="utf-8", errors="replace") as f:
        return [(0, clean_text(f.read()))]


def _extract_image(file_path: str) -> list[tuple[int, str]]:
    """
    Send an image to a vision LLM and return its text description.
    The description is what gets embedded and stored — makes images searchable.
    """
    with open(file_path, "rb") as f:
        image_data = base64.b64encode(f.read()).decode("utf-8")

    ext = Path(file_path).suffix.lower().lstrip(".")
    mime_map = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "webp": "webp", "gif": "gif"}
    mime_type = f"image/{mime_map.get(ext, 'jpeg')}"

    response = _openrouter.chat.completions.create(
        model=settings.vision_model,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime_type};base64,{image_data}"},
                    },
                    {
                        "type": "text",
                        "text": (
                            "Describe this image in full detail. Include: any visible text, "
                            "numbers, labels, titles; the content of charts, graphs or diagrams; "
                            "objects, people, settings; and any other relevant information. "
                            "Write in plain prose so the description is searchable."
                        ),
                    },
                ],
            }
        ],
        max_tokens=1024,
    )

    description = response.choices[0].message.content
    filename = Path(file_path).name
    tagged = f"[Image: {filename}]\n{description}"
    return [(0, tagged)]


def _extract_pages(file_path: str) -> list[tuple[int, str]]:
    """Route to the correct extractor based on file extension."""
    ext = Path(file_path).suffix.lower()
    if ext in PDF_EXTENSIONS:
        return _extract_pdf(file_path)
    if ext in DOCX_EXTENSIONS:
        return _extract_docx(file_path)
    if ext in TEXT_EXTENSIONS:
        return _extract_txt(file_path)
    if ext in IMAGE_EXTENSIONS:
        return _extract_image(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


# ── Pinecone helpers ──────────────────────────────────────────────────────────

def _embed_batch(texts: list[str]) -> list[list[float]]:
    """Embed a list of texts using Pinecone Inference API, in batches of 96."""
    pc = get_pc()
    vectors: list[list[float]] = []
    batch_size = 96
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        response = pc.inference.embed(
            model=settings.embedding_model,
            inputs=batch,
            parameters={"input_type": "passage", "truncate": "END"},
        )
        vectors.extend([item["values"] for item in response.data])
    return vectors


def clear_index(index_name: str | None = None) -> None:
    """Delete ALL vectors from the given index (keeps the index itself)."""
    get_index(index_name).delete(delete_all=True)


def delete_document(filename: str, index_name: str | None = None) -> int:
    """
    Delete all Pinecone vectors that belong to *filename*.
    Returns the number of vectors deleted (approximate via stats diff).
    """
    index = get_index(index_name)
    before = index.describe_index_stats().total_vector_count
    index.delete(filter={"source_file": {"$eq": filename}})
    after = index.describe_index_stats().total_vector_count
    return max(0, before - after)


# ── Main entry point ──────────────────────────────────────────────────────────

def ingest(file_path: str, index_name: str | None = None) -> list[dict]:
    """
    Parse, chunk, embed, and upsert a document into Pinecone.
    Supports: PDF, DOCX, TXT, PNG, JPG, JPEG, WEBP, GIF.
    Returns a list of chunk metadata dicts.
    """
    filename = Path(file_path).name
    pages = _extract_pages(file_path)

    # Build flat list of chunk records preserving page provenance
    chunks: list[dict] = []
    chunk_index = 0
    for page_num, page_text in pages:
        splits = _splitter.split_text(page_text)
        for text in splits:
            chunks.append({
                "chunk_index": chunk_index,
                "source_file": filename,
                "page": page_num,
                "text": text,
            })
            chunk_index += 1

    if not chunks:
        return []

    # Embed all chunks
    texts = [c["text"] for c in chunks]
    vectors = _embed_batch(texts)

    # Build Pinecone upsert payload
    records = [
        {
            "id": make_vector_id(filename, c["chunk_index"]),
            "values": vec,
            "metadata": {
                "source_file": c["source_file"],
                "page": c["page"],
                "chunk_index": c["chunk_index"],
                "text": c["text"],
            },
        }
        for c, vec in zip(chunks, vectors)
    ]

    index = get_index(index_name)
    batch_size = 100
    for i in range(0, len(records), batch_size):
        index.upsert(vectors=records[i : i + batch_size])

    return chunks
